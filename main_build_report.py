# -*- coding: utf-8 -*-
r"""
main_build_report.py
Orchestrates the creation of the complete academic project report:
"D:\uhv\PestRisk_UHV_Project_Report.docx"
matching the exact formatting, styling, sections, and tables of
"D:\safecall-java\SafeCall_AI_Project_Report.docx"
tailored for Universal Human Values (UHV II), Life Skills and Community Work.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from report_front_matter import (
    build_cover_page,
    build_declaration_and_certificate,
    build_preliminary_matter
)
from report_chapters import build_section3_chapters

def main():
    print("="*60)
    print("Initiating build of PestRisk UHV Project Report (.docx)...")
    print("="*60)

    doc = docx.Document()

    # Set normal style default font to Times New Roman 12 pt
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Times New Roman'
    font_normal.size = Pt(12)
    font_normal.color.rgb = RGBColor(0, 0, 0)

    # 1. Build Section 0: Cover / Title Page
    print("Building Section 0: Title / Cover Page...")
    build_cover_page(doc)

    # 2. Build Section 1: Declaration & Certificate
    print("Building Section 1: Declaration & Certificate...")
    build_declaration_and_certificate(doc)

    # 3. Build Section 2: Preliminary Matter (Acknowledgement, Vision/Mission, POs, COs, Abstract, TOC, List of Tables, Abbreviations)
    print("Building Section 2: Preliminary Matter...")
    build_preliminary_matter(doc)

    # 4. Build Section 3: Chapters 1 to 7 and References
    print("Building Section 3: Chapters 1 to 7 & References...")
    build_section3_chapters(doc)

    # 5. Save Document
    candidates = [
        os.path.join(r"D:\uhv", "PestRisk_UHV_Project_Report.docx"),
        os.path.join(r"D:\uhv", "PestRisk_UHV_Project_Report_v2.docx"),
        os.path.join(r"D:\uhv", "PestRisk_UHV_Project_Report_v3.docx"),
        os.path.join(r"D:\uhv", "PestRisk_UHV_Project_Report_Updated.docx")
    ]
    
    saved_path = None
    for cand in candidates:
        try:
            doc.save(cand)
            saved_path = cand
            print(f"SUCCESS: Report saved successfully to:\n{cand}")
            break
        except PermissionError:
            print(f"NOTE: {cand} is currently open/locked in Word. Trying next filename...")
    
    if not saved_path:
        raise PermissionError("All candidate report files are locked open in Word.")

    # Validate output file properties
    file_size = os.path.getsize(saved_path)
    doc_check = docx.Document(saved_path)
    print("\n--- VALIDATION SUMMARY ---")
    print(f"Target File: {saved_path}")
    print(f"File Size: {file_size:,} bytes")
    print(f"Total Paragraphs: {len(doc_check.paragraphs)}")
    print(f"Total Tables: {len(doc_check.tables)}")
    print(f"Total Sections: {len(doc_check.sections)}")
    print(f"Total Inline Shapes: {len(doc_check.inline_shapes)}")
    print("="*60)

if __name__ == "__main__":
    main()
