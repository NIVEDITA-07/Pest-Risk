# -*- coding: utf-8 -*-
"""
build_pestrisk_full_report.py
Generates the complete academic project report:
"PestRisk_UHV_Project_Report.docx"
for Course 24PWT208: UHV II, LIFE SKILLS AND COMMUNITY WORK
Department of Computer Science and Engineering
Sahrdaya College of Engineering and Technology (Autonomous), Kodakara, Thrissur
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def add_header_text(header, text="PestRisk"):
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(9.5)
    r.bold = True
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

def add_footer_page_number(footer, prefix="SCET, Kodakara\t\t", initial_val="1"):
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = p.add_run(prefix)
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(9)
    r1.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    t_run = OxmlElement('w:r')
    t_text = OxmlElement('w:t')
    t_text.text = initial_val
    t_run.append(t_text)
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')

    r_field = p.add_run()
    r_field._r.append(fldChar1)
    r_field._r.append(instrText)
    r_field._r.append(fldChar2)
    r_field._r.append(t_run)
    r_field._r.append(fldChar3)
    r_field.font.name = "Times New Roman"
    r_field.font.size = Pt(10)
    r_field.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

def add_p(doc, text="", font_size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(font_size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return p

def add_bullet(doc, title, text, font_size=12, space_before=0, space_after=4, line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    r_bullet = p.add_run("•  ")
    r_bullet.font.name = "Times New Roman"
    r_bullet.font.size = Pt(font_size)
    r_bullet.bold = True
    
    if title:
        r_title = p.add_run(title)
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(font_size)
        r_title.bold = True
        
    r_text = p.add_run(text)
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(font_size)
    r_text.bold = False
    return p

def add_num_item(doc, num_str, title, text, font_size=12, space_before=2, space_after=5, line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    r_num = p.add_run(f"{num_str}  ")
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(font_size)
    r_num.bold = True
    
    if title:
        r_title = p.add_run(title)
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(font_size)
        r_title.bold = True
        
    r_text = p.add_run(text)
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(font_size)
    r_text.bold = False
    return p

def add_chapter_title(doc, chapter_num, title_text):
    p1 = add_p(doc, f"CHAPTER {chapter_num}", font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=4)
    p2 = add_p(doc, title_text.upper(), font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=18)
    return p1, p2

def add_sec_title(doc, sec_num, title_text):
    return add_p(doc, f"{sec_num}  {title_text.upper()}", font_size=13, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=14, space_after=6)

def add_subsec_title(doc, subsec_num, title_text):
    return add_p(doc, f"{subsec_num}  {title_text}", font_size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=4)

def add_table_caption(doc, caption_text):
    return add_p(doc, f"Table: {caption_text}", font_size=10, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=4)
